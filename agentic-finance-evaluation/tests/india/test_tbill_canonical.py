"""RBI 91-Day T-Bill primary auction yield weekly acquisition/validation tests.

Covers the RBI '50 Macroeconomic Indicators' workbook Weekly sheet series
'91-Day Treasury Bill (Primary) Yield (%)'. The raw workbook is immutable
evidence (shared with the 10Y milestone) and is never modified by tests.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest

from src.india.tbill_canonicalize import (
    TENOR_SERIES,
    build_canonical,
    detect_overlaps,
    discover_tbill_raw_files,
    expected_tbill_header,
    header_matches_tbill_schema,
    parse_period,
    read_weekly_rows,
    validate_tbill_file,
)
from src.india.leakage_audit import LeakageAuditor
from src.india.manifest import ManifestManager
from src.schemas.india_data import (
    AcquisitionStatus,
    DataFrequency,
    EligibilityStatus,
    ValidationStatus,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_DIR = PROJECT_ROOT / "data" / "raw" / "india" / "fixed_income"
CANONICAL_PATH = PROJECT_ROOT / "data" / "processed" / "india" / "market" / "rbi_tbill_91d_weekly.csv"
MANIFEST_ID = "rbi_gsec_91d_yield"
RAW_NAME = "50 Macroeconomic Indicators.xlsx"

CANCELLED_WEEKS = ["2026-03-27", "2025-02-21"]
PLACEHOLDER_WEEKS = [
    "2026-04-03", "2026-03-27", "2026-01-02", "2025-02-21",
    "2024-09-27", "2024-09-20", "2023-03-31",
]


def _workbook_with_series(tmp_path: Path, name: str, entries: list) -> Path:
    import openpyxl

    path = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly"
    ws.append([])
    ws.append([])
    ws.append([])
    ws.append(["", "Period"] + [""] * 7 + ["91-Day Treasury Bill (Primary) Yield (%)"])
    for day, value in entries:
        ws.append(["", day, "", "", "", "", "", "", "", value])
    wb.save(path)
    return path


# 1. raw discovery (shared workbook artifact)
def test_raw_workbook_discovered():
    files = discover_tbill_raw_files(RAW_DIR)
    assert [path.name for path in files] == [RAW_NAME]


# 2. exact source/series targeting (column J, 91D only)
def test_series_targeting():
    series, column = expected_tbill_header("91D")
    assert series == "91-Day Treasury Bill (Primary) Yield (%)"
    assert column == 9
    assert "364D" in TENOR_SERIES
    with pytest.raises(ValueError):
        expected_tbill_header("NOPE")
    header, _ = read_weekly_rows(RAW_DIR / RAW_NAME)
    assert header_matches_tbill_schema(header, "91D") is True


# 3. schema
def test_schema():
    audit, frame, excluded = validate_tbill_file(RAW_DIR / RAW_NAME, RAW_NAME, "91D")
    assert audit.header_valid is True
    assert audit.series == "91-Day Treasury Bill (Primary) Yield (%)"
    assert audit.row_count == 465
    assert audit.parsed_date_count == 465
    assert set(frame.columns) >= {"observation_date", "tenor", "yield_pct"}


# 4. date parsing (Friday Periods)
def test_date_parsing():
    assert parse_period(datetime(2026, 9, 4)) == date(2026, 9, 4)
    with pytest.raises(ValueError):
        parse_period("NOT-A-DATE")
    with pytest.raises(ValueError):
        parse_period(None)
    audit, _, _ = validate_tbill_file(RAW_DIR / RAW_NAME, RAW_NAME, "91D")
    assert audit.all_fridays is True
    assert audit.min_date == date(2017, 10, 13)
    assert audit.max_date == date(2026, 9, 4)


# 5. numeric validity
def test_numeric_validity():
    audit, _, _ = validate_tbill_file(RAW_DIR / RAW_NAME, RAW_NAME, "91D")
    assert audit.numeric_parsing_failures == 0
    assert audit.null_or_malformed_rows == 0
    assert audit.nonpositive_count == 0
    df = pd.read_csv(CANONICAL_PATH)
    assert (df["yield_pct"] > 0).all()


# 6. duplicate handling (none in artifact)
def test_duplicate_handling():
    audit, _, _ = validate_tbill_file(RAW_DIR / RAW_NAME, RAW_NAME, "91D")
    assert audit.duplicate_dates_within_file == 0
    df = pd.read_csv(CANONICAL_PATH)
    assert df["observation_date"].duplicated().sum() == 0


# 7. event-frequency semantics (weekly events; cancelled weeks stay absent)
def test_event_frequency_semantics():
    audit, frame, excluded = validate_tbill_file(RAW_DIR / RAW_NAME, RAW_NAME, "91D")
    assert audit.placeholder_count == 7
    assert sorted(e.date_text for e in excluded) == sorted(PLACEHOLDER_WEEKS)
    df = pd.read_csv(CANONICAL_PATH)
    present = set(df["observation_date"])
    for week in PLACEHOLDER_WEEKS:
        assert week not in present
    # Cancelled-auction weeks: present in raw as '-' placeholders, absent in canonical
    for week in CANCELLED_WEEKS:
        assert week not in present
    assert len(df) == 458


# 8. chronological ordering (raw descending, canonical ascending)
def test_chronological_ordering():
    audit, _, _ = validate_tbill_file(RAW_DIR / RAW_NAME, RAW_NAME, "91D")
    assert audit.is_descending is True
    df = pd.read_csv(CANONICAL_PATH)
    dates = pd.to_datetime(df["observation_date"])
    assert dates.is_monotonic_increasing
    assert df["observation_date"].iloc[0] == "2017-10-13"
    assert df["observation_date"].iloc[-1] == "2026-09-04"


# 9. canonical uniqueness + bookkeeping
def test_canonical_uniqueness():
    df = pd.read_csv(CANONICAL_PATH)
    assert len(df) == 458
    assert df["observation_date"].nunique() == 458
    assert (df["tenor"] == "91D").all()


# 10. provenance (shared raw artifact with 10Y milestone)
def test_provenance():
    df = pd.read_csv(CANONICAL_PATH)
    assert "source_files" in df.columns
    assert (df["source"] == "RBI").all()
    assert (df["source_files"] == RAW_NAME).all()
    manager = ManifestManager(str(PROJECT_ROOT))
    manifest = manager.load_manifest(MANIFEST_ID)
    assert manifest.raw_paths == [f"data/raw/india/fixed_income/{RAW_NAME}"]
    assert manifest.raw_artifacts is not None and len(manifest.raw_artifacts) == 1
    assert manifest.raw_artifacts[0].sha256 == manager.compute_sha256(
        manifest.raw_artifacts[0].path
    )
    assert manifest.frequency == DataFrequency.WEEKLY


# 11. SHA reproducibility
def test_sha_reproducibility():
    manager = ManifestManager(str(PROJECT_ROOT))
    manifest = manager.load_manifest(MANIFEST_ID)
    assert manifest.processed_sha256 == manager.compute_sha256(
        "data/processed/india/market/rbi_tbill_91d_weekly.csv"
    )
    assert manifest.raw_artifacts
    for artifact in manifest.raw_artifacts:
        assert artifact.sha256 == manager.compute_sha256(artifact.path)


# 12. overlap/conflict rejection (single source: none; synthetic conflict rejected)
def test_overlap_conflict_rejection(tmp_path: Path):
    files = discover_tbill_raw_files(RAW_DIR)
    frames = [validate_tbill_file(p, p.name, "91D")[1] for p in files]
    assert detect_overlaps(frames, [p.name for p in files]) == []

    import openpyxl

    def _wb(path: Path, day: datetime, value: float) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Weekly"
        ws.append([])
        ws.append([])
        ws.append([])
        ws.append(["", "Period"] + [""] * 7 + ["91-Day Treasury Bill (Primary) Yield (%)"])
        ws.append(["", day, "", "", "", "", "", "", "", value])
        wb.save(path)

    first = tmp_path / "a.xlsx"
    second = tmp_path / "b.xlsx"
    _wb(first, datetime(2020, 1, 3), 5.0)
    _wb(second, datetime(2020, 1, 3), 9.9)
    frames2 = [
        validate_tbill_file(first, first.name, "91D")[1],
        validate_tbill_file(second, second.name, "91D")[1],
    ]
    overlaps = detect_overlaps(frames2, [first.name, second.name])
    assert len(overlaps) == 1 and overlaps[0].identical is False
    with pytest.raises(ValueError, match="Conflicting overlapping"):
        build_canonical(frames2, [first.name, second.name], overlaps)


# 13. leakage/information boundary
def test_leakage_information_boundary():
    df = pd.read_csv(CANONICAL_PATH)
    assert not any(
        any(token in str(col).lower() for token in ("future", "forward", "next_", "lead_"))
        for col in df.columns
    )
    report = LeakageAuditor().audit_all(price_df=df)
    assert report.is_clean
    cutoff = pd.Timestamp("2022-01-01")
    visible = df[pd.to_datetime(df["observation_date"]) <= cutoff]
    hidden = df[pd.to_datetime(df["observation_date"]) > cutoff]
    assert len(visible) > 0 and len(hidden) > 0
    assert pd.to_datetime(visible["observation_date"]).max() <= cutoff
    assert pd.to_datetime(hidden["observation_date"]).min() > cutoff


# 14. blocked eligibility
def test_blocked_eligibility():
    manager = ManifestManager(str(PROJECT_ROOT))
    manifest = manager.load_manifest(MANIFEST_ID)
    assert manifest.acquisition_status == AcquisitionStatus.ACQUIRED
    assert manifest.validation_status == ValidationStatus.PASSED
    assert manifest.eligibility_status == EligibilityStatus.BLOCKED
    assert manifest.eligibility_status != EligibilityStatus.EXPERIMENT_ELIGIBLE


# 15. manifest bookkeeping (465 - 7 - 0 = 458)
def test_manifest_bookkeeping():
    manager = ManifestManager(str(PROJECT_ROOT))
    manifest = manager.load_manifest(MANIFEST_ID)
    params = manifest.processing_parameters
    assert params["total_raw_rows"] == 465
    assert len(params["excluded_invalid_rows"]) == 7
    assert params["duplicate_extra_rows_before_dedup"] == 0
    assert manifest.row_count == 458 == manifest.unique_date_count


# 16. no fabrication of missing auction observations
def test_no_fabrication():
    df = pd.read_csv(CANONICAL_PATH)
    dates = pd.to_datetime(df["observation_date"]).dt.date
    gaps = sorted((b - a).days for a, b in zip(sorted(dates)[:-1], sorted(dates)[1:]))
    assert 7 in gaps
    assert max(gaps) > 7
    assert len(df) < 465 + 1
