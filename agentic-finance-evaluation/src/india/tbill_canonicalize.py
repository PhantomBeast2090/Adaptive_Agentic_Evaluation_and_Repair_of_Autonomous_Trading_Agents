"""RBI Treasury Bill primary auction yield weekly canonicalization.

Validates the RBI "50 Macroeconomic Indicators" workbook in place (raw
file is never rewritten) and builds canonical T-Bill datasets under
``data/processed/india/market/``.

Raw evidence: the workbook's Weekly sheet. Supported series (column
positions verified against the header row):

    91D  -> "91-Day Treasury Bill (Primary) Yield (%)"   (column J)
    364D -> "364-Day Treasury Bill (Primary) Yield (%)"  (column L)

Values are primary-market auction cut-off (implicit YTM) yields in %
p.a. for week-ending Friday rows. RBI marks weeks without a published
observation with a ``-`` placeholder (e.g. cancelled-auction weeks such
as 2026-03-27 and 2025-02-21); placeholders are EXCLUDED with documented
reasons and never filled, interpolated, or forward-filled.

Canonical schema follows ``IndianFixedIncomeRecord``:
    observation_date, tenor, yield_pct, source, source_files

No publication_date is supplied in the workbook, so none is recorded;
the agent information boundary is the observation (week-ending) date.
Weekly event-frequency observations are never converted to daily.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

WORKBOOK_NAME = "50 Macroeconomic Indicators.xlsx"
TARGET_SHEET = "Weekly"
PERIOD_COLUMN = 1
TENOR_SERIES: Dict[str, Tuple[str, int]] = {
    "91D": ("91-Day Treasury Bill (Primary) Yield (%)", 9),
    "364D": ("364-Day Treasury Bill (Primary) Yield (%)", 11),
}
CANONICAL_COLUMNS = [
    "observation_date",
    "tenor",
    "yield_pct",
    "source",
    "source_files",
]


def compute_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def discover_tbill_raw_files(raw_dir: Path) -> List[Path]:
    """Return the RBI macro workbook(s). Raw files are never modified."""
    return sorted(raw_dir.glob("50 Macroeconomic Indicators.xlsx"))


def read_weekly_rows(path: Path) -> Tuple[List[str], List[List[Any]]]:
    """Read Weekly sheet header + data rows via openpyxl (read-only)."""
    import openpyxl
    import warnings

    warnings.filterwarnings("ignore")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if TARGET_SHEET not in workbook.sheetnames:
        raise ValueError(f"{path.name}: sheet {TARGET_SHEET!r} not found")
    sheet = workbook[TARGET_SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    header = [("" if value is None else str(value).strip()) for value in rows[3]]
    return header, [list(row) for row in rows[4:]]


def expected_tbill_header(tenor: str) -> Tuple[str, int]:
    """Return (series name, column index) for a tenor. Raises if unknown."""
    if tenor not in TENOR_SERIES:
        raise ValueError(f"Unknown tenor {tenor!r}. Supported: {sorted(TENOR_SERIES)}")
    return TENOR_SERIES[tenor]


def header_matches_tbill_schema(normalized: List[str], tenor: str) -> bool:
    """Check Period column and exact target series column for a tenor."""
    series, column = expected_tbill_header(tenor)
    return (
        len(normalized) > column
        and normalized[PERIOD_COLUMN] == "Period"
        and normalized[column] == series
    )


def parse_period(value: Any) -> date:
    """Parse RBI weekly Period (datetime Friday). Raises on malformed."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"non-date Period: {value!r}")


@dataclass
class TbillExcludedRow:
    source_file: str
    tenor: str
    date_text: str
    reason: str


@dataclass
class TbillFileAudit:
    filename: str
    rel_path: str
    tenor: str
    byte_size: int
    sha256: str
    sheet: str
    series: str
    header_valid: bool
    row_count: int
    parsed_date_count: int
    valid_row_count: int
    excluded_row_count: int
    min_date: Optional[date]
    max_date: Optional[date]
    all_fridays: bool
    duplicate_dates_within_file: int
    null_or_malformed_rows: int
    numeric_parsing_failures: int
    placeholder_count: int
    nonpositive_count: int
    is_descending: bool
    is_ascending: bool


@dataclass
class TbillOverlap:
    overlap_date: date
    files: List[str]
    identical: bool
    records: List[Dict[str, Any]] = field(default_factory=list)


def validate_tbill_file(
    path: Path, rel_path: str, tenor: str = "91D"
) -> Tuple[TbillFileAudit, pd.DataFrame, List[TbillExcludedRow]]:
    """Validate the workbook's target tenor series in place.

    Returns (audit, valid-records frame, excluded rows with reasons).
    """
    series, column = expected_tbill_header(tenor)
    raw_bytes = path.read_bytes()
    sha = hashlib.sha256(raw_bytes).hexdigest()
    header, data_rows = read_weekly_rows(path)
    header_valid = header_matches_tbill_schema(header, tenor)

    records: List[Dict[str, Any]] = []
    excluded: List[TbillExcludedRow] = []
    null_malformed = 0
    numeric_failures = 0
    placeholders = 0
    nonpositive = 0
    parsed_dates: List[date] = []
    seen: Dict[date, int] = {}

    for row in data_rows:
        if len(row) <= column:
            null_malformed += 1
            continue
        period_raw, yield_raw = row[PERIOD_COLUMN], row[column]
        if period_raw is None or (isinstance(period_raw, str) and not period_raw.strip()):
            null_malformed += 1
            continue
        try:
            obs = parse_period(period_raw)
        except ValueError:
            null_malformed += 1
            continue
        if yield_raw is None or (isinstance(yield_raw, str) and yield_raw.strip() in ("", "-")):
            placeholders += 1
            parsed_dates.append(obs)
            seen[obs] = seen.get(obs, 0) + 1
            excluded.append(
                TbillExcludedRow(
                    source_file=f"{path.name}#{TARGET_SHEET}",
                    tenor=tenor,
                    date_text=obs.isoformat(),
                    reason="RBI non-observation placeholder ('-' in source; no auction observation published)",
                )
            )
            continue
        try:
            yield_val = float(yield_raw)
        except (TypeError, ValueError):
            numeric_failures += 1
            continue
        parsed_dates.append(obs)
        seen[obs] = seen.get(obs, 0) + 1
        if not yield_val > 0:
            nonpositive += 1
            excluded.append(
                TbillExcludedRow(
                    source_file=f"{path.name}#{TARGET_SHEET}",
                    tenor=tenor,
                    date_text=obs.isoformat(),
                    reason="non-positive yield (% p.a. must be > 0)",
                )
            )
            continue
        records.append(
            {
                "observation_date": obs,
                "tenor": tenor,
                "yield_pct": yield_val,
                "source_file": path.name,
            }
        )

    dup_within = sum(count - 1 for count in seen.values() if count > 1)
    is_desc = (
        all(parsed_dates[i] >= parsed_dates[i + 1] for i in range(len(parsed_dates) - 1))
        if len(parsed_dates) > 1
        else True
    )
    is_asc = (
        all(parsed_dates[i] <= parsed_dates[i + 1] for i in range(len(parsed_dates) - 1))
        if len(parsed_dates) > 1
        else True
    )
    audit = TbillFileAudit(
        filename=path.name,
        rel_path=rel_path,
        tenor=tenor,
        byte_size=len(raw_bytes),
        sha256=sha,
        sheet=TARGET_SHEET,
        series=series,
        header_valid=header_valid,
        row_count=len(data_rows),
        parsed_date_count=len(parsed_dates),
        valid_row_count=len(records),
        excluded_row_count=len(excluded),
        min_date=min(parsed_dates) if parsed_dates else None,
        max_date=max(parsed_dates) if parsed_dates else None,
        all_fridays=all(d.weekday() == 4 for d in seen) if seen else False,
        duplicate_dates_within_file=dup_within,
        null_or_malformed_rows=null_malformed,
        numeric_parsing_failures=numeric_failures,
        placeholder_count=placeholders,
        nonpositive_count=nonpositive,
        is_descending=is_desc,
        is_ascending=is_asc,
    )
    frame = pd.DataFrame(records)
    return audit, frame, excluded


def audit_all_files(
    raw_dir: Path, tenor: str = "91D"
) -> Tuple[List[TbillFileAudit], List[pd.DataFrame], List[TbillExcludedRow]]:
    files = discover_tbill_raw_files(raw_dir)
    audits: List[TbillFileAudit] = []
    frames: List[pd.DataFrame] = []
    excluded_all: List[TbillExcludedRow] = []
    for path in files:
        rel = f"data/raw/india/fixed_income/{path.name}"
        audit, frame, excluded = validate_tbill_file(path, rel, tenor)
        audits.append(audit)
        frames.append(frame)
        excluded_all.extend(excluded)
    return audits, frames, excluded_all


def detect_overlaps(
    frames: List[pd.DataFrame], filenames: List[str]
) -> List[TbillOverlap]:
    """Group valid records by date across files; flag identical vs conflicting."""
    by_date: Dict[date, List[Tuple[str, float]]] = {}
    for frame, name in zip(frames, filenames):
        for _, row in frame.iterrows():
            key = (
                row["observation_date"]
                if isinstance(row["observation_date"], date)
                else pd.to_datetime(row["observation_date"]).date()
            )
            by_date.setdefault(key, []).append((name, float(row["yield_pct"])))
    overlaps: List[TbillOverlap] = []
    for obs_date in sorted(by_date):
        entries = by_date[obs_date]
        if len(entries) <= 1:
            continue
        first_value = entries[0][1]
        identical = all(value == first_value for _, value in entries)
        overlaps.append(
            TbillOverlap(
                overlap_date=obs_date,
                files=[name for name, _ in entries],
                identical=identical,
                records=[
                    {"source_file": name, "values": [value]}
                    for name, value in entries
                ],
            )
        )
    return overlaps


def build_canonical(
    frames: List[pd.DataFrame],
    filenames: List[str],
    overlaps: List[TbillOverlap],
) -> pd.DataFrame:
    """Combine validated frames, dedup only proven identical dates.

    Raises ValueError if any conflicting overlap exists — the caller must
    report the conflict rather than choosing a record arbitrarily.
    """
    conflicts = [item for item in overlaps if not item.identical]
    if conflicts:
        detail = "; ".join(
            f"{item.overlap_date} in {item.files}" for item in conflicts[:5]
        )
        raise ValueError(f"Conflicting overlapping records detected: {detail}")
    combined = pd.concat(frames, ignore_index=True)
    # Provenance: every raw file that carried this observation date.
    provenance: Dict[date, List[str]] = {}
    for frame, name in zip(frames, filenames):
        for value in frame["observation_date"].tolist():
            key = value if isinstance(value, date) else pd.to_datetime(value).date()
            provenance.setdefault(key, [])
            if name not in provenance[key]:
                provenance[key].append(name)
    # Deduplicate identical observations: keep first occurrence, sorted
    # ascending by date, never forward-fill or fabricate.
    combined["_sort_date"] = pd.to_datetime(combined["observation_date"])
    combined = combined.sort_values("_sort_date").drop_duplicates(
        subset=["observation_date"], keep="first"
    )
    combined["source_files"] = combined["observation_date"].map(
        lambda d: ";".join(
            sorted(
                provenance[d if isinstance(d, date) else pd.to_datetime(d).date()]
            )
        )
    )
    combined["observation_date"] = pd.to_datetime(combined["observation_date"]).dt.date
    combined["source"] = "RBI"
    combined = combined.sort_values("observation_date").reset_index(drop=True)
    return combined[CANONICAL_COLUMNS]
