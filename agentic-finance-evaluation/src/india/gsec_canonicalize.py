"""RBI 10Y G-Sec yield (FBIL) weekly canonicalization.

Validates the RBI "50 Macroeconomic Indicators" workbook in place (raw
file is never rewritten) and builds a canonical dataset under
``data/processed/india/market/``.

Raw evidence: a single ``.xlsx`` workbook with sheets Weekly /
Fortnightly / Monthly / Quarterly. Only the Weekly sheet's exact series
``10-Year G-Sec Yield (FBIL) (%)`` (column M, Period in column B) is the
target dataset; no other sheet or column is used. Periods are week-ending
Fridays; values are yields in % p.a. as republished by RBI from the FBIL
G-Sec valuation benchmarks (cubic-spline curve; see FBIL G-Sec Valuation
Methodology v4). The canonical variable is therefore the RBI-published
weekly 10Y G-Sec (FBIL-derived) yield — a fixed-income market-state
variable, not an auction cut-off and not a single-security price.

Canonical schema follows ``IndianFixedIncomeRecord``:
    observation_date, tenor, yield_pct, source, source_files

No publication_date is supplied in the workbook, so none is recorded;
the agent information boundary is the observation (week-ending) date.
Weekly observations are never converted to daily, forward-filled, or
interpolated. Pre-2017 history is not spliced from any other source.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

WORKBOOK_NAME = "50 Macroeconomic Indicators.xlsx"
TARGET_SHEET = "Weekly"
TARGET_SERIES = "10-Year G-Sec Yield (FBIL) (%)"
TENOR = "10Y"
CANONICAL_COLUMNS = [
    "observation_date",
    "tenor",
    "yield_pct",
    "source",
    "source_files",
]


def compute_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def discover_gsec_raw_files(raw_dir: Path) -> List[Path]:
    """Return the RBI macro workbook(s). Raw files are never modified."""
    return sorted(raw_dir.glob("50 Macroeconomic Indicators.xlsx"))


def read_weekly_series(path: Path) -> Tuple[List[str], List[List[Any]]]:
    """Read Weekly sheet header + data rows via openpyxl (read-only)."""
    import openpyxl
    import warnings

    warnings.filterwarnings("ignore")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if TARGET_SHEET not in workbook.sheetnames:
        raise ValueError(f"{path.name}: sheet {TARGET_SHEET!r} not found")
    sheet = workbook[TARGET_SHEET]
    rows = list(sheet.iter_rows(values_only=True))
    header = [("" if value is None else str(value).strip()) for value in rows[3]]
    return header, [list(row) for row in rows[4:]]


def header_matches_gsec_schema(normalized: List[str]) -> bool:
    """Check Period column and exact target series column positions."""
    return (
        len(normalized) > 12
        and normalized[1] == "Period"
        and normalized[12] == TARGET_SERIES
    )


def parse_period(value: Any) -> date:
    """Parse RBI weekly Period (datetime Friday). Raises on malformed."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"non-date Period: {value!r}")


@dataclass
class GsecExcludedRow:
    source_file: str
    date_text: str
    reason: str


@dataclass
class GsecFileAudit:
    filename: str
    rel_path: str
    byte_size: int
    sha256: str
    sheet: str
    series: str
    header_valid: bool
    row_count: int
    parsed_date_count: int
    valid_row_count: int
    excluded_row_count: int
    min_date: Optional[date]
    max_date: Optional[date]
    all_fridays: bool
    duplicate_dates_within_file: int
    null_or_malformed_rows: int
    numeric_parsing_failures: int
    nonpositive_count: int
    is_descending: bool
    is_ascending: bool


@dataclass
class GsecOverlap:
    overlap_date: date
    files: List[str]
    identical: bool
    records: List[Dict[str, Any]] = field(default_factory=list)


def validate_gsec_file(
    path: Path, rel_path: str
) -> Tuple[GsecFileAudit, pd.DataFrame, List[GsecExcludedRow]]:
    """Validate the workbook's target series in place.

    Returns (audit, valid-records frame, excluded rows with reasons).
    """
    raw_bytes = path.read_bytes()
    sha = hashlib.sha256(raw_bytes).hexdigest()
    header, data_rows = read_weekly_series(path)
    header_valid = header_matches_gsec_schema(header)

    records: List[Dict[str, Any]] = []
    excluded: List[GsecExcludedRow] = []
    null_malformed = 0
    numeric_failures = 0
    nonpositive = 0
    parsed_dates: List[date] = []
    seen: Dict[date, int] = {}

    for row in data_rows:
        if len(row) < 13:
            null_malformed += 1
            continue
        period_raw, yield_raw = row[1], row[12]
        if period_raw is None or (isinstance(period_raw, str) and not period_raw.strip()):
            null_malformed += 1
            continue
        try:
            obs = parse_period(period_raw)
        except ValueError:
            null_malformed += 1
            continue
        if yield_raw is None or (isinstance(yield_raw, str) and yield_raw.strip() in ("", "-")):
            null_malformed += 1
            continue
        try:
            yield_val = float(yield_raw)
        except (TypeError, ValueError):
            numeric_failures += 1
            continue
        parsed_dates.append(obs)
        seen[obs] = seen.get(obs, 0) + 1
        if not yield_val > 0:
            nonpositive += 1
            excluded.append(
                GsecExcludedRow(
                    source_file=f"{path.name}#{TARGET_SHEET}",
                    date_text=obs.isoformat(),
                    reason="non-positive yield (% p.a. must be > 0)",
                )
            )
            continue
        records.append(
            {
                "observation_date": obs,
                "tenor": TENOR,
                "yield_pct": yield_val,
                "source_file": path.name,
            }
        )

    dup_within = sum(count - 1 for count in seen.values() if count > 1)
    is_desc = (
        all(parsed_dates[i] >= parsed_dates[i + 1] for i in range(len(parsed_dates) - 1))
        if len(parsed_dates) > 1
        else True
    )
    is_asc = (
        all(parsed_dates[i] <= parsed_dates[i + 1] for i in range(len(parsed_dates) - 1))
        if len(parsed_dates) > 1
        else True
    )
    audit = GsecFileAudit(
        filename=path.name,
        rel_path=rel_path,
        byte_size=len(raw_bytes),
        sha256=sha,
        sheet=TARGET_SHEET,
        series=TARGET_SERIES,
        header_valid=header_valid,
        row_count=len(data_rows),
        parsed_date_count=len(parsed_dates),
        valid_row_count=len(records),
        excluded_row_count=len(excluded),
        min_date=min(parsed_dates) if parsed_dates else None,
        max_date=max(parsed_dates) if parsed_dates else None,
        all_fridays=all(d.weekday() == 4 for d in seen) if seen else False,
        duplicate_dates_within_file=dup_within,
        null_or_malformed_rows=null_malformed,
        numeric_parsing_failures=numeric_failures,
        nonpositive_count=nonpositive,
        is_descending=is_desc,
        is_ascending=is_asc,
    )
    frame = pd.DataFrame(records)
    return audit, frame, excluded


def audit_all_files(
    raw_dir: Path,
) -> Tuple[List[GsecFileAudit], List[pd.DataFrame], List[GsecExcludedRow]]:
    files = discover_gsec_raw_files(raw_dir)
    audits: List[GsecFileAudit] = []
    frames: List[pd.DataFrame] = []
    excluded_all: List[GsecExcludedRow] = []
    for path in files:
        rel = f"data/raw/india/fixed_income/{path.name}"
        audit, frame, excluded = validate_gsec_file(path, rel)
        audits.append(audit)
        frames.append(frame)
        excluded_all.extend(excluded)
    return audits, frames, excluded_all


def detect_overlaps(
    frames: List[pd.DataFrame], filenames: List[str]
) -> List[GsecOverlap]:
    """Group valid records by date across files; flag identical vs conflicting."""
    by_date: Dict[date, List[Tuple[str, float]]] = {}
    for frame, name in zip(frames, filenames):
        for _, row in frame.iterrows():
            key = (
                row["observation_date"]
                if isinstance(row["observation_date"], date)
                else pd.to_datetime(row["observation_date"]).date()
            )
            by_date.setdefault(key, []).append((name, float(row["yield_pct"])))
    overlaps: List[GsecOverlap] = []
    for obs_date in sorted(by_date):
        entries = by_date[obs_date]
        if len(entries) <= 1:
            continue
        first_value = entries[0][1]
        identical = all(value == first_value for _, value in entries)
        overlaps.append(
            GsecOverlap(
                overlap_date=obs_date,
                files=[name for name, _ in entries],
                identical=identical,
                records=[
                    {"source_file": name, "values": [value]}
                    for name, value in entries
                ],
            )
        )
    return overlaps


def build_canonical(
    frames: List[pd.DataFrame],
    filenames: List[str],
    overlaps: List[GsecOverlap],
) -> pd.DataFrame:
    """Combine validated frames, dedup only proven identical dates.

    Raises ValueError if any conflicting overlap exists — the caller must
    report the conflict rather than choosing a record arbitrarily.
    """
    conflicts = [item for item in overlaps if not item.identical]
    if conflicts:
        detail = "; ".join(
            f"{item.overlap_date} in {item.files}" for item in conflicts[:5]
        )
        raise ValueError(f"Conflicting overlapping records detected: {detail}")
    combined = pd.concat(frames, ignore_index=True)
    # Provenance: every raw file that carried this observation date.
    provenance: Dict[date, List[str]] = {}
    for frame, name in zip(frames, filenames):
        for value in frame["observation_date"].tolist():
            key = value if isinstance(value, date) else pd.to_datetime(value).date()
            provenance.setdefault(key, [])
            if name not in provenance[key]:
                provenance[key].append(name)
    # Deduplicate identical observations: keep first occurrence, sorted
    # ascending by date, never forward-fill or fabricate.
    combined["_sort_date"] = pd.to_datetime(combined["observation_date"])
    combined = combined.sort_values("_sort_date").drop_duplicates(
        subset=["observation_date"], keep="first"
    )
    combined["source_files"] = combined["observation_date"].map(
        lambda d: ";".join(
            sorted(
                provenance[d if isinstance(d, date) else pd.to_datetime(d).date()]
            )
        )
    )
    combined["observation_date"] = pd.to_datetime(combined["observation_date"]).dt.date
    combined["source"] = "RBI"
    combined = combined.sort_values("observation_date").reset_index(drop=True)
    return combined[CANONICAL_COLUMNS]
